"""Reviewer decisions (ADR-011): policy, recording, supersede rules, the
effective view, gate preconditions, trail digests, exports and the API.

The artifacts come from DEMO_MODE (fixed, schema-valid content through the
real AuditFlow), so every rule is exercised on a realistic session:
CTRL-01 (key, ToD Effective / ToE Not tested → No exception), CTRL-02 (key,
ToD Ineffective → Exception, DEF-01 proposed Medium) and CTRL-03 (key, Not
tested).
"""

import io
import os
import sys
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from review_helpers import (  # type: ignore[import-not-found]
    record_demo_decisions,
    record_flow_decisions,
)
from api.exports import ExportContext, report_markdown, working_papers_xlsx
from api.job_store import remove_flow, set_flow
from swarm import review_policy as policy
from swarm import session_manager
from swarm import trail as audit_trail
from swarm.audit_flow import (
    AuditFlow,
    DecisionConflictError,
    DecisionValidationError,
    MissingReviewDecisionsError,
    SegregationOfDutiesError,
)
from swarm.demo import demo_review_decisions
from swarm.review_decisions import DecisionContext, effective_view
from swarm.state.repository import FlowRepository

AUTH = {"Authorization": "Bearer test-token"}
PREPARER = "Pat Preparer"
IN_CHARGE = "Ivan In-Charge"
MANAGER = "Mona Manager"


@pytest.fixture(autouse=True)
def demo_env(monkeypatch):
    monkeypatch.setenv("DEMO_MODE", "1")
    monkeypatch.setenv("DEMO_STEP_DELAY", "0")
    monkeypatch.delenv("ENVIRONMENT", raising=False)
    monkeypatch.delenv("DEMO_QA_REJECT_PHASE", raising=False)
    monkeypatch.delenv("VAULT_ENCRYPTION_KEY", raising=False)


def _at_gate_2(required: bool = True, theme: str = "S3 exposure") -> AuditFlow:
    flow = AuditFlow()
    flow.state.theme = theme
    flow.state.business_context = "Fintech storing customer data"
    flow.record_preparer(PREPARER, require_review_decisions=required)
    flow.begin_phase_1()
    flow.generate_planning()
    flow.begin_phase_2(IN_CHARGE)
    flow.generate_fieldwork()
    assert flow.state.status == "WAITING_HUMAN_GATE_2"
    return flow


def _at_gate_3(theme: str = "S3 exposure") -> AuditFlow:
    flow = _at_gate_2(theme=theme)
    record_flow_decisions(flow, "gate2", IN_CHARGE)
    flow.begin_phase_3(IN_CHARGE)
    flow.generate_reporting()
    assert flow.state.status == "WAITING_HUMAN_GATE_3"
    return flow


def _completed() -> AuditFlow:
    flow = _at_gate_3()
    record_flow_decisions(flow, "gate3", MANAGER)
    flow.finalize_audit(MANAGER)
    assert flow.state.status == "COMPLETED"
    return flow


def _decide(flow: AuditFlow, decision_type: str, subject_id: str, **kw):
    kw.setdefault(
        "decided_by", IN_CHARGE if flow.state.status.endswith("2") else MANAGER
    )
    return flow.record_decision(
        decision_type=decision_type, subject_id=subject_id, **kw
    )


RESPONSE = {
    "text": "Agreed; BPA will be enabled.",
    "agreement": "agree",
    "action_owner_role": "Cloud Platform Lead",
    "target_date": "2026-12-31",
    "received_from": "CISO",
    "received_on": "2026-09-20",
}


# ── Policy functions ─────────────────────────────────────────────────────────


class TestPolicy:
    @pytest.mark.parametrize(
        "draft,new,allowed",
        [
            ("Effective", "Not tested", True),
            ("Effective", "Effective", True),
            ("Ineffective", "Not tested", False),
            ("Exceptions noted", "Not tested", False),
            ("Effective", "Ineffective", False),
            ("Not tested", "Effective", False),
        ],
    )
    def test_conclusion_change_allowed(self, draft, new, allowed):
        assert policy.conclusion_change_allowed(draft, new) is allowed

    @pytest.mark.parametrize(
        "result,key,needed",
        [
            ("Exception", False, True),
            ("No exception", True, True),
            ("Not tested", True, True),
            ("No exception", False, False),
            ("Not tested", None, False),
        ],
    )
    def test_finding_needs_review(self, result, key, needed):
        assert policy.finding_needs_review(result, key) is needed

    def test_scope_limitation_only_for_untested_key_controls(self):
        assert policy.finding_needs_scope_limitation("Not tested", True)
        assert not policy.finding_needs_scope_limitation("Not tested", False)
        assert not policy.finding_needs_scope_limitation("Exception", True)

    @pytest.mark.parametrize("dtype", sorted(policy.PREPARER_EXCLUDED_DECISIONS))
    def test_preparer_excluded(self, dtype):
        assert policy.decision_sod_violation(dtype, " pat  PREPARER", PREPARER)
        assert policy.decision_sod_violation(dtype, IN_CHARGE, PREPARER) is None

    @pytest.mark.parametrize("dtype", ["writeup", "management_response"])
    def test_preparer_may_draft_and_transcribe(self, dtype):
        assert policy.decision_sod_violation(dtype, PREPARER, PREPARER) is None

    def test_status_rules(self):
        assert (
            policy.decision_status_violation("sign_off", "WAITING_HUMAN_GATE_2") is None
        )
        assert policy.decision_status_violation("sign_off", "WAITING_HUMAN_GATE_3")
        assert policy.decision_status_violation("classify", "WAITING_HUMAN_GATE_2")
        assert policy.decision_status_violation("classify", "COMPLETED")
        assert (
            policy.decision_status_violation("management_response", "COMPLETED") is None
        )
        assert policy.decision_status_violation("nonsense", "COMPLETED")

    def test_classification_differs(self):
        draft = {"classification": "Medium", "likelihood": "High", "magnitude": "Low"}
        assert not policy.classification_differs(draft, dict(draft))
        assert policy.classification_differs(draft, {**draft, "magnitude": "High"})

    def test_change_rate_is_not_published(self):
        assert policy.REVIEWER_CHANGE_RATE_PUBLISHED is False


# ── Recording: who, when, what ───────────────────────────────────────────────


class TestRecording:
    def test_preparer_cannot_sign_off(self):
        flow = _at_gate_2()
        with pytest.raises(SegregationOfDutiesError):
            _decide(flow, "sign_off", "CTRL-01", decided_by=PREPARER)
        assert flow.state.review_decisions == []

    def test_preparer_can_draft_writeup_and_transcribe_response(self):
        flow = _at_gate_3()
        values = {
            k: f"{k} text"
            for k in ("criteria", "condition", "cause", "effect", "recommendation")
        }
        _decide(flow, "writeup", "DEF-01", decided_by=PREPARER, values=values)
        _decide(
            flow, "management_response", "DEF-01", decided_by=PREPARER, values=RESPONSE
        )
        # ... but the write-up does not stand in for the reviewer's classification.
        missing = flow.effective_view().missing_for_gate["3"]
        assert any(m.required == ["classify"] for m in missing)

    def test_decision_is_append_only_and_drafts_unchanged(self):
        flow = _at_gate_2()
        before = audit_trail.artifact_digest(flow.state.working_papers)
        d = _decide(
            flow,
            "challenge",
            "CTRL-01",
            values={"tod_conclusion": "Not tested"},
            rationale="The read predates the period.",
        )
        assert audit_trail.artifact_digest(flow.state.working_papers) == before
        assert d.draft_digest == before
        assert d.phase == 2 and d.artifact == "working_papers"
        assert d.identity_source == "declared"
        entry = flow.state.approval_trail[-1]
        assert entry["action"] == "review_decision"
        assert entry["decision_id"] == d.decision_id
        assert entry["decision_digest"] == audit_trail.decision_digest(d)
        assert entry["phase"] == "2"
        assert entry["subject"] == "finding:CTRL-01"

    def test_phase_2_decisions_sealed_after_gate_2(self):
        flow = _at_gate_3()
        with pytest.raises(DecisionConflictError, match="WAITING_HUMAN_GATE_2"):
            _decide(flow, "sign_off", "CTRL-01")

    def test_classify_not_before_report(self):
        flow = _at_gate_2()
        with pytest.raises(DecisionConflictError):
            _decide(flow, "classify", "DEF-01", values={"classification": "High"})

    def test_after_completed_only_management_responses(self):
        flow = _completed()
        with pytest.raises(DecisionConflictError):
            _decide(
                flow,
                "classify",
                "DEF-01",
                values={"classification": "Low"},
                rationale="x",
            )
        d = _decide(
            flow, "management_response", "DEF-01", decided_by=PREPARER, values=RESPONSE
        )
        assert d.phase == 3
        assert flow.verify_trail()["status"] == "ok"

    @pytest.mark.parametrize(
        "kwargs,match",
        [
            ({"decision_type": "approve", "subject_id": "CTRL-01"}, "decision_type"),
            (
                {"decision_type": "sign_off", "subject_id": "CTRL-99"},
                "no working-paper",
            ),
            ({"decision_type": "sign_off", "subject_id": ""}, "subject_id"),
            (
                {
                    "decision_type": "sign_off",
                    "subject_id": "CTRL-01",
                    "subject_type": "deficiency",
                },
                "about a finding",
            ),
            ({"decision_type": "challenge", "subject_id": "CTRL-01"}, "rationale"),
            (
                {
                    "decision_type": "sign_off",
                    "subject_id": "CTRL-01",
                    "values": {"classification": "High"},
                },
                "values",
            ),
            (
                {
                    "decision_type": "sign_off",
                    "subject_id": "CTRL-01",
                    "decided_by": " ",
                },
                "decided_by",
            ),
        ],
    )
    def test_validation_errors_at_gate_2(self, kwargs, match):
        flow = _at_gate_2()
        kwargs.setdefault("decided_by", IN_CHARGE)
        with pytest.raises(DecisionValidationError, match=match):
            flow.record_decision(**kwargs)
        assert flow.state.review_decisions == []
        assert flow.state.approval_trail[-1]["action"] != "review_decision"

    def test_challenge_may_withdraw_a_positive_conclusion(self):
        flow = _at_gate_2()
        _decide(
            flow,
            "challenge",
            "CTRL-01",
            values={"tod_conclusion": "Not tested", "toe_conclusion": "Not tested"},
            rationale="Configuration read is outside the period.",
        )
        fv = flow.effective_view().finding("CTRL-01")
        assert fv is not None
        assert fv.review_status == "challenged"
        assert fv.draft.result == "No exception"
        assert fv.effective.tod_conclusion == "Not tested"
        assert fv.effective.result == "Not tested"
        assert fv.differs_from_draft

    @pytest.mark.parametrize(
        "control,values",
        [
            ("CTRL-02", {"tod_conclusion": "Not tested"}),  # removes an exception
            ("CTRL-01", {"tod_conclusion": "Ineffective"}),  # new conclusion
            ("CTRL-03", {"toe_conclusion": "Effective"}),
        ],
    )
    def test_other_conclusion_changes_need_rework(self, control, values):
        flow = _at_gate_2()
        with pytest.raises(DecisionConflictError, match="rework"):
            _decide(flow, "challenge", control, values=values, rationale="r")

    def test_challenge_without_change_is_a_recorded_reservation(self):
        flow = _at_gate_2()
        d = _decide(flow, "challenge", "CTRL-02", rationale="Sample of three is thin.")
        assert d.values == {}
        fv = flow.effective_view().finding("CTRL-02")
        assert fv is not None and fv.review_status == "challenged"
        assert not fv.differs_from_draft

    def test_classify_same_as_draft_needs_no_rationale(self):
        flow = _at_gate_3()
        d = _decide(flow, "classify", "DEF-01", values={"classification": "Medium"})
        # Likelihood and magnitude default to the draft's.
        assert d.values == {
            "classification": "Medium",
            "likelihood": "High",
            "magnitude": "Medium",
        }

    def test_classify_different_needs_rationale(self):
        flow = _at_gate_3()
        with pytest.raises(DecisionValidationError, match="rationale"):
            _decide(flow, "classify", "DEF-01", values={"classification": "High"})
        with pytest.raises(DecisionValidationError, match="rationale"):
            _decide(
                flow,
                "classify",
                "DEF-01",
                values={"classification": "Medium", "magnitude": "High"},
            )
        _decide(
            flow, "classify", "DEF-01", values={"classification": "High"}, rationale="r"
        )

    def test_classify_must_use_the_report_scale(self):
        flow = _at_gate_3()
        with pytest.raises(DecisionValidationError, match="scale"):
            _decide(
                flow,
                "classify",
                "DEF-01",
                values={"classification": "Significant Deficiency"},
                rationale="r",
            )

    def test_material_weakness_rule(self):
        flow = _at_gate_3(theme="SOX ITGC financial reporting")
        with pytest.raises(DecisionValidationError, match="Material Weakness"):
            _decide(
                flow,
                "classify",
                "DEF-01",
                values={"classification": "Material Weakness", "likelihood": "Low"},
                rationale="r",
            )
        _decide(
            flow,
            "classify",
            "DEF-01",
            values={
                "classification": "Material Weakness",
                "likelihood": "Medium",
                "magnitude": "High",
            },
            rationale="Aggregated with the IAM exceptions.",
        )

    def test_scope_limitation_only_for_untested_findings(self):
        flow = _at_gate_3()
        with pytest.raises(DecisionValidationError, match="Not tested"):
            _decide(flow, "scope_limitation", "CTRL-02", rationale="r")
        with pytest.raises(DecisionValidationError, match="rationale"):
            _decide(flow, "scope_limitation", "CTRL-03")
        _decide(flow, "scope_limitation", "CTRL-03", rationale="Held outside AWS.")

    def test_engagement_conclusion_rules(self):
        flow = _at_gate_3()
        with pytest.raises(DecisionValidationError, match="values"):
            _decide(
                flow,
                "engagement_conclusion",
                "",
                values={"conclusion": "Great"},
                rationale="r",
            )
        with pytest.raises(DecisionValidationError, match="rationale"):
            _decide(
                flow, "engagement_conclusion", "", values={"conclusion": "Satisfactory"}
            )
        with pytest.raises(DecisionValidationError, match="engagement"):
            _decide(
                flow,
                "engagement_conclusion",
                "DEF-01",
                values={"conclusion": "Satisfactory"},
                rationale="r",
            )
        d = _decide(
            flow,
            "engagement_conclusion",
            "",
            values={"conclusion": "needs improvement"},
            rationale="One high deficiency.",
        )
        assert d.subject_id == "engagement"
        assert d.values == {"conclusion": "Needs improvement"}

    def test_writeup_needs_all_five_parts(self):
        flow = _at_gate_3()
        with pytest.raises(DecisionValidationError, match="cause"):
            _decide(
                flow,
                "writeup",
                "DEF-01",
                values={
                    "criteria": "c",
                    "condition": "c",
                    "effect": "e",
                    "recommendation": "r",
                },
            )

    @pytest.mark.parametrize(
        "change,match",
        [
            ({"action_owner_role": None}, "action_owner_role"),
            ({"target_date": "end of year"}, "ISO date"),
            ({"received_on": ""}, "received_on"),
            ({"agreement": "maybe"}, "agreement"),
            ({"agreement": "disagree"}, "rebuttal"),
            ({"received_on": "2099-01-01"}, "after the date"),
        ],
    )
    def test_management_response_validation(self, change, match):
        flow = _at_gate_3()
        values = {k: v for k, v in {**RESPONSE, **change}.items() if v is not None}
        with pytest.raises(DecisionValidationError, match=match):
            _decide(flow, "management_response", "DEF-01", values=values)

    def test_management_disagreement_with_rebuttal(self):
        flow = _at_gate_3()
        d = _decide(
            flow,
            "management_response",
            "DEF-01",
            values={
                "text": "The bucket is intentionally public.",
                "agreement": "disagree",
                "received_from": "CISO",
                "received_on": "2026-09-20",
            },
            rationale="It holds customer exports; public access is not justified.",
        )
        assert "action_owner_role" not in d.values
        md = report_markdown(
            flow.state.final_report,
            flow.state.approval_trail,
            ExportContext("sid", "Demo", flow.state.status),
            flow.effective_view(),
        )
        assert "management disagrees" in md
        assert "**Auditor's rebuttal:** It holds customer exports" in md


# ── Supersede and stale decisions ────────────────────────────────────────────


class TestSupersede:
    def test_second_decision_must_supersede(self):
        flow = _at_gate_2()
        first = _decide(flow, "sign_off", "CTRL-02")
        with pytest.raises(DecisionConflictError, match=first.decision_id):
            _decide(flow, "challenge", "CTRL-02", rationale="r")

    def test_supersede_replaces_the_active_decision(self):
        flow = _at_gate_2()
        first = _decide(flow, "sign_off", "CTRL-02")
        second = _decide(
            flow,
            "challenge",
            "CTRL-02",
            rationale="Sample too small",
            supersedes=first.decision_id,
        )
        states = flow.decision_states()
        assert states == {first.decision_id: "superseded", second.decision_id: "active"}
        view = flow.effective_view()
        fv = view.finding("CTRL-02")
        assert fv is not None and fv.review_status == "challenged"
        assert view.superseded_decision_ids == [first.decision_id]
        assert flow.state.approval_trail[-1]["supersedes"] == first.decision_id

    def test_cannot_supersede_twice_or_across_subjects(self):
        flow = _at_gate_2()
        first = _decide(flow, "sign_off", "CTRL-02")
        _decide(flow, "sign_off", "CTRL-02", supersedes=first.decision_id)
        with pytest.raises(DecisionConflictError, match="superseded"):
            _decide(flow, "sign_off", "CTRL-02", supersedes=first.decision_id)
        other = _decide(flow, "sign_off", "CTRL-01")
        with pytest.raises(DecisionValidationError, match="same subject"):
            _decide(flow, "sign_off", "CTRL-02", supersedes=other.decision_id)
        with pytest.raises(DecisionValidationError, match="no decision"):
            _decide(flow, "sign_off", "CTRL-03", supersedes="nope")

    def test_nothing_to_supersede(self):
        flow = _at_gate_2()
        a = _decide(flow, "sign_off", "CTRL-01")
        flow2 = _at_gate_2()
        flow2.state.review_decisions.append(a)  # decision from another draft
        with pytest.raises(DecisionConflictError, match="stale"):
            _decide(flow2, "sign_off", "CTRL-01", supersedes=a.decision_id)

    def test_rework_makes_decisions_stale(self):
        flow = _at_gate_2()
        record_flow_decisions(flow, "gate2", IN_CHARGE)
        assert flow.effective_view().missing_for_gate["2"] == []
        flow.return_for_rework(2, MANAGER, "Re-perform CTRL-02 over the period.")
        flow.generate_fieldwork()
        view = flow.effective_view()
        assert len(view.stale_decision_ids) == 3
        assert {m.subject_id for m in view.missing_for_gate["2"]} == {
            "CTRL-01",
            "CTRL-02",
            "CTRL-03",
        }
        # A fresh decision on the new draft needs no supersedes.
        _decide(flow, "sign_off", "CTRL-02")
        with pytest.raises(MissingReviewDecisionsError):
            flow.begin_phase_3(IN_CHARGE)


# ── Gate preconditions ───────────────────────────────────────────────────────


class TestGatePreconditions:
    def test_gate_2_lists_missing_sign_offs(self):
        flow = _at_gate_2()
        with pytest.raises(MissingReviewDecisionsError) as exc:
            flow.begin_phase_3(IN_CHARGE)
        assert {m["subject_id"] for m in exc.value.missing} == {
            "CTRL-01",
            "CTRL-02",
            "CTRL-03",
        }
        ctrl2 = next(m for m in exc.value.missing if m["subject_id"] == "CTRL-02")
        assert ctrl2["required"] == ["sign_off", "challenge"]
        assert "Exception" in ctrl2["reason"]
        assert flow.state.status == "WAITING_HUMAN_GATE_2"
        assert flow.state.approval_trail[-1]["gate"] == "Gate 1 (Planning)"

    def test_gate_3_lists_classify_scope_and_conclusion(self):
        flow = _at_gate_3()
        with pytest.raises(MissingReviewDecisionsError) as exc:
            flow.finalize_audit(MANAGER)
        assert sorted(
            (m["subject_type"], m["subject_id"]) for m in exc.value.missing
        ) == [
            ("deficiency", "DEF-01"),
            ("engagement", "engagement"),
            ("finding", "CTRL-03"),
        ]

    def test_sod_is_still_checked_first(self):
        flow = _at_gate_3()
        record_flow_decisions(flow, "gate3", MANAGER)
        with pytest.raises(SegregationOfDutiesError):
            flow.finalize_audit(IN_CHARGE)  # approved Gate 2

    def test_legacy_audit_needs_no_decisions(self):
        flow = _at_gate_2(required=False)
        assert not flow.review_decisions_required()
        flow.begin_phase_3(IN_CHARGE)
        flow.generate_reporting()
        flow.finalize_audit(MANAGER)
        assert flow.state.status == "COMPLETED"

    def test_requirement_survives_editing_the_stored_flag(self):
        flow = _at_gate_2()
        flow.state.review_decisions_required = False
        assert flow.review_decisions_required()  # the chained trail entry says so
        with pytest.raises(MissingReviewDecisionsError):
            flow.begin_phase_3(IN_CHARGE)

    def test_approval_seals_the_phase_decisions(self):
        flow = _completed()
        gates = [e for e in flow.state.approval_trail if e["action"] == "gate_approval"]
        g2, g3 = gates[1], gates[2]
        phase2 = [d for d in flow.state.review_decisions if d.phase == 2]
        assert g2["decisions_count"] == str(len(phase2)) == "3"
        assert g2["decisions_digest"] == audit_trail.decisions_digest(phase2)
        assert int(g3["decisions_count"]) == 4  # classify, writeup, scope, conclusion


# ── Trail verification ───────────────────────────────────────────────────────


class TestTrailDetection:
    def test_completed_walkthrough_verifies(self):
        flow = _completed()
        result = flow.verify_trail()
        assert result["ok"], result
        assert result["decisions_changed"] == []

    def test_edited_decision_after_approval_is_detected(self):
        flow = _completed()
        decisions = flow.state.review_decisions
        i = next(k for k, d in enumerate(decisions) if d.decision_type == "classify")
        decisions[i] = decisions[i].model_copy(
            update={"values": {**decisions[i].values, "classification": "Low"}}
        )
        result = flow.verify_trail()
        assert result["status"] == audit_trail.STATUS_ARTIFACT_CHANGED
        assert result["changed_since_approval"] == ["Gate 3 (Reporting)"]
        assert "changed after it was recorded" in result["decisions_changed"][0]

    def test_edited_decision_before_approval_is_detected(self):
        flow = _at_gate_2()
        d = _decide(flow, "sign_off", "CTRL-01")
        flow.state.review_decisions[0] = d.model_copy(update={"decided_by": "Someone"})
        result = flow.verify_trail()
        assert result["status"] == audit_trail.STATUS_DECISION_CHANGED
        assert not result["ok"]

    def test_removed_and_injected_decisions_are_detected(self):
        flow = _at_gate_2()
        d = _decide(flow, "sign_off", "CTRL-01")
        flow.state.review_decisions.clear()
        assert "removed" in flow.verify_trail()["decisions_changed"][0]
        flow.state.review_decisions.append(d.model_copy(update={"decision_id": "x"}))
        problems = flow.verify_trail()["decisions_changed"]
        assert any("no trail entry" in p for p in problems)

    def test_removed_sealed_decision_changes_the_gate(self):
        flow = _completed()
        flow.state.review_decisions.pop(0)  # a Gate 2 sign-off
        result = flow.verify_trail()
        assert "Gate 2 (Fieldwork)" in result["changed_since_approval"]

    def test_legacy_trail_without_decisions_still_verifies(self):
        """A trail written before decisions existed: gate approvals carry no
        decisions_digest and there are no review_decision entries."""
        flow = _at_gate_2(required=False)
        papers = flow.state.working_papers
        trail: list[dict[str, str]] = []
        audit_trail.append_entry(
            trail, {"gate": "Audit created", "action": "audit_created"}
        )
        audit_trail.append_entry(
            trail,
            {
                "gate": "Gate 2 (Fieldwork)",
                "action": "gate_approval",
                "artifact": "working_papers",
                "artifact_digest": audit_trail.artifact_digest(papers),
            },
        )
        result = audit_trail.verify_trail(
            trail, artifacts={"working_papers": papers}, decisions=[]
        )
        assert result["ok"], result
        assert audit_trail.verify_trail(trail, artifacts={"working_papers": papers})[
            "ok"
        ]


# ── Effective view and rendering ─────────────────────────────────────────────


class TestEffectiveViewAndExports:
    def _walked(self) -> AuditFlow:
        flow = _at_gate_2()
        _decide(flow, "sign_off", "CTRL-02")
        _decide(flow, "sign_off", "CTRL-03")
        _decide(
            flow,
            "challenge",
            "CTRL-01",
            values={"tod_conclusion": "Not tested"},
            rationale="Read outside the period.",
        )
        flow.begin_phase_3(IN_CHARGE)
        flow.generate_reporting()
        return flow

    def test_reviewer_change_rate(self):
        flow = self._walked()
        _decide(
            flow, "classify", "DEF-01", values={"classification": "High"}, rationale="r"
        )
        rate = flow.effective_view().reviewer_change_rate
        assert (rate.subjects_decided, rate.subjects_changed) == (4, 2)
        assert rate.rate == pytest.approx(0.5)
        assert rate.published is False

    def test_empty_change_rate(self):
        flow = _at_gate_2()
        rate = flow.effective_view().reviewer_change_rate
        assert rate.subjects_decided == 0 and rate.rate is None

    def test_view_from_context(self):
        flow = self._walked()
        ctx = DecisionContext(
            flow.state.racm_plan, flow.state.working_papers, flow.state.final_report
        )
        view = effective_view(ctx, flow.state.review_decisions)
        dv = view.deficiency("DEF-01")
        assert dv is not None and dv.classification_source == "ai_draft"
        assert view.deficiency_scale == "Risk rating"

    def test_report_shows_conclusion_of_record_and_ai_draft(self):
        flow = self._walked()
        record_flow_decisions(flow, "gate3", MANAGER)
        flow.finalize_audit(MANAGER)
        ctx = ExportContext("sid", "Demo", flow.state.status)
        assert flow.state.final_report is not None
        md = report_markdown(
            flow.state.final_report,
            flow.state.approval_trail,
            ctx,
            flow.effective_view(),
        )
        assert "## Deficiency Evaluation (conclusion of record)" in md
        assert "Deficiency Evaluation (proposed)" not in md
        assert "High — reviewer: Mona Manager (declared identity)" in md
        assert "AI draft: Medium" in md
        assert "## Engagement Conclusion" in md and "**Needs improvement**" in md
        assert "- **Criteria:**" in md
        assert "| CTRL-01 | Not tested (AI draft: No exception" in md
        assert "| Challenged | Ivan In-Charge" in md
        assert "## Scope Limitations" in md and "**CTRL-03**" in md
        assert "(review_decision) — classify on deficiency:DEF-01" in md
        assert "conclusion of record is the reviewer's decision" in md
        # A short block near the top states the reviewer's classification and
        # engagement conclusion, and labels the AI-drafted summary as such.
        assert "## Conclusions of Record" in md
        assert md.index("## Conclusions of Record") < md.index("## Executive Summary")
        assert "## Executive Summary (AI-drafted summary)" in md
        assert "**Engagement conclusion:** Needs improvement" in md
        assert "DEF-01: High (AI draft: Medium) — decided by" in md
        assert "**Scope limitations:**" in md and "CTRL-03 — decided by" in md
        # Without the view (legacy callers) the proposed wording is unchanged.
        legacy = report_markdown(flow.state.final_report, [], ctx)
        assert "## Deficiency Evaluation (proposed)" in legacy
        assert "## Conclusions of Record" not in legacy
        assert "## Executive Summary\n" in legacy

    def test_report_omits_conclusions_of_record_before_any_decision(self):
        flow = self._walked()
        md = report_markdown(
            flow.state.final_report,
            [],
            ExportContext("sid", "Demo", flow.state.status),
            flow.effective_view(),
        )
        assert "## Conclusions of Record" not in md
        assert "## Executive Summary\n" in md
        assert "AI-drafted summary" not in md

    def test_report_marks_undecided_rows_as_proposed(self):
        flow = self._walked()
        report = flow.state.final_report
        assert report is not None
        second = report.deficiency_evaluations[0].model_copy(
            update={"deficiency_id": "DEF-02", "title": "Second"}
        )
        flow.state.final_report = report.model_copy(
            update={"deficiency_evaluations": [*report.deficiency_evaluations, second]}
        )
        _decide(flow, "classify", "DEF-01", values={"classification": "Medium"})
        md = report_markdown(
            flow.state.final_report,
            [],
            ExportContext("sid", "Demo", flow.state.status),
            flow.effective_view(),
        )
        assert "(conclusion of record where decided)" in md
        assert "| Medium (proposed) |" in md
        assert "Medium — reviewer: Mona Manager" in md

    def test_working_papers_xlsx_shows_review(self):
        flow = self._walked()
        assert flow.state.working_papers is not None
        body = working_papers_xlsx(
            flow.state.working_papers,
            ExportContext("sid", "Demo", flow.state.status),
            flow.effective_view(),
        )
        ws = load_workbook(io.BytesIO(body))["Findings"]
        rows = list(ws.iter_rows(values_only=True))
        header = rows[0]
        by_control = {r[0]: dict(zip(header, r)) for r in rows[1:]}
        ctrl1 = by_control["CTRL-01"]
        assert ctrl1["ToD Conclusion"] == "Not tested"
        assert ctrl1["Result"] == "Not tested"
        assert ctrl1["Reviewer Decision"] == "Challenged"
        assert ctrl1["Reviewed By (declared)"] == IN_CHARGE
        assert "ToD Effective" in ctrl1["AI Draft (where the record differs)"]
        assert by_control["CTRL-02"]["Reviewer Decision"] == "Signed off"
        assert by_control["CTRL-02"]["AI Draft (where the record differs)"] is None


# ── Demo walk-through helper ─────────────────────────────────────────────────


def test_demo_review_decisions_rejects_unknown_stage():
    with pytest.raises(ValueError):
        demo_review_decisions("gate1", {})


def test_demo_review_decisions_are_complete_for_icfr():
    flow = _at_gate_3(theme="SOX ITGC financial reporting")
    record_flow_decisions(flow, "gate3", MANAGER)
    flow.finalize_audit(MANAGER)
    dv = flow.effective_view().deficiency("DEF-01")
    assert dv is not None and not dv.differs_from_draft


# ── API ──────────────────────────────────────────────────────────────────────


class _InlineExecutor:
    def submit(self, session_id, fn, *args):
        fn(*args)


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("API_AUTH_TOKEN", "test-token")
    monkeypatch.setenv("TRAIL_ANCHORS_PATH", str(tmp_path / "anchors.json"))
    monkeypatch.setattr(
        session_manager, "SESSIONS_PATH", str(tmp_path / "audit_sessions.json")
    )
    from api.main import app

    with patch("api.routers.sessions.get_executor", return_value=_InlineExecutor()):
        yield TestClient(app)


def _new_session(c) -> str:
    r = c.post(
        "/api/sessions",
        headers=AUTH,
        json={
            "theme": "S3 exposure",
            "business_context": "Fintech",
            "prepared_by": PREPARER,
        },
    )
    assert r.status_code == 201
    sid = r.json()["session_id"]
    r = c.patch(
        f"/api/sessions/{sid}/approve",
        headers=AUTH,
        json={"gate_number": 1, "human_id": IN_CHARGE},
    )
    assert r.status_code == 200
    return sid


def _post(c, sid, **body):
    body.setdefault("decided_by", IN_CHARGE)
    return c.post(f"/api/sessions/{sid}/decisions", headers=AUTH, json=body)


def test_api_decision_lifecycle(client):
    c = client
    sid = _new_session(c)
    try:
        detail = c.get(f"/api/sessions/{sid}", headers=AUTH).json()
        assert detail["review_decisions_required"] is True
        assert len(detail["effective"]["missing_for_gate"]["2"]) == 3

        r = c.patch(
            f"/api/sessions/{sid}/approve",
            headers=AUTH,
            json={"gate_number": 2, "human_id": IN_CHARGE},
        )
        assert r.status_code == 409
        body = r.json()
        assert isinstance(body["detail"], str) and "CTRL-02" in body["detail"]
        assert {m["subject_id"] for m in body["missing_decisions"]} == {
            "CTRL-01",
            "CTRL-02",
            "CTRL-03",
        }

        # 409: preparer; 409: wrong phase; 422: validation.
        assert (
            _post(
                c,
                sid,
                decision_type="sign_off",
                subject_id="CTRL-01",
                decided_by=PREPARER,
            ).status_code
            == 409
        )
        assert (
            _post(
                c,
                sid,
                decision_type="classify",
                subject_id="DEF-01",
                values={"classification": "High"},
            ).status_code
            == 409
        )
        assert (
            _post(c, sid, decision_type="sign_off", subject_id="CTRL-42").status_code
            == 422
        )
        assert (
            _post(
                c, sid, decision_type="sign_off", subject_id="CTRL-01", values={"x": 1}
            ).status_code
            == 422
        )
        assert (
            _post(c, sid, decision_type="bogus", subject_id="CTRL-01").status_code
            == 422
        )
        assert (
            c.post(
                "/api/sessions/nope/decisions",
                headers=AUTH,
                json={"decision_type": "sign_off", "decided_by": "x"},
            ).status_code
            == 404
        )

        r = _post(c, sid, decision_type="sign_off", subject_id="CTRL-02")
        assert r.status_code == 201, r.text
        first = r.json()
        assert first["state"] == "active" and first["identity_source"] == "declared"
        assert (
            _post(c, sid, decision_type="sign_off", subject_id="CTRL-02").status_code
            == 409
        )
        r = _post(
            c,
            sid,
            decision_type="challenge",
            subject_id="CTRL-02",
            rationale="Sample too small",
            supersedes=first["decision_id"],
        )
        assert r.status_code == 201, r.text

        listing = c.get(f"/api/sessions/{sid}/decisions", headers=AUTH).json()
        assert [d["state"] for d in listing["decisions"]] == ["superseded", "active"]
        assert listing["effective"]["findings"][2]["review_status"] == "challenged"

        record_demo_decisions(c, sid, "gate2", IN_CHARGE, AUTH)
        r = c.patch(
            f"/api/sessions/{sid}/approve",
            headers=AUTH,
            json={"gate_number": 2, "human_id": IN_CHARGE},
        )
        assert r.status_code == 200, r.text
        record_demo_decisions(c, sid, "gate3", MANAGER, AUTH)
        r = c.patch(
            f"/api/sessions/{sid}/approve",
            headers=AUTH,
            json={"gate_number": 3, "human_id": MANAGER},
        )
        assert r.status_code == 200, r.text
        record_demo_decisions(c, sid, "after_issue", PREPARER, AUTH)

        # Served from disk: decisions, states and the effective view survive.
        remove_flow(sid)
        detail = c.get(f"/api/sessions/{sid}", headers=AUTH).json()
        assert detail["status"] == "COMPLETED"
        assert detail["trail_verification"]["ok"] is True, detail["trail_verification"]
        dv = detail["effective"]["deficiencies"][0]
        assert dv["classification_source"] == "reviewer"
        assert dv["effective"]["classification"] == "High"
        assert dv["draft"]["classification"] == "Medium"
        assert dv["management_response"]["values"]["agreement"] == "agree"
        assert {d["state"] for d in detail["review_decisions"]} == {
            "active",
            "superseded",
        }
        listing = c.get(f"/api/sessions/{sid}/decisions", headers=AUTH).json()
        assert len(listing["decisions"]) == len(detail["review_decisions"])

        report = c.get(f"/api/sessions/{sid}/export/report.md", headers=AUTH).text
        assert "(conclusion of record)" in report
        assert "Management response" in report
        verify = c.get(f"/api/sessions/{sid}/trail/verify", headers=AUTH).json()
        assert verify["ok"] is True and verify["decisions_changed"] == []
    finally:
        remove_flow(sid)


def test_api_legacy_snapshot_without_decisions_loads(client):
    c = client
    sid = "legacy-session"
    flow = AuditFlow(initial_status="WAITING_HUMAN_GATE_2")
    session_manager.save_session(sid, "legacy", "ctx")
    FlowRepository().save(sid, flow)
    # Strip the new fields, as an older snapshot would lack them.
    data = session_manager.get_session(sid)
    assert data is not None
    snap = data["state_snapshot"]
    snap.pop("review_decisions")
    snap.pop("review_decisions_required")
    session_manager.update_session(sid, state_snapshot=snap)
    remove_flow(sid)
    detail = c.get(f"/api/sessions/{sid}", headers=AUTH).json()
    assert detail["review_decisions"] == []
    assert detail["review_decisions_required"] is False
    assert detail["effective"]["findings"] == []
    loaded = FlowRepository().load(sid)
    assert loaded is not None and loaded.is_clean
    assert loaded.flow.state.review_decisions == []
    set_flow(sid, loaded.flow)
    try:
        assert c.get(f"/api/sessions/{sid}/decisions", headers=AUTH).json() == {
            "decisions": [],
            "effective": detail["effective"],
        }
    finally:
        remove_flow(sid)


def test_walkthrough_script_regenerates_the_sample(monkeypatch, tmp_path):
    """scripts/demo_walkthrough.py drives the API end to end and exports."""
    monkeypatch.setenv("API_AUTH_TOKEN", "test-token")
    monkeypatch.setenv("ENVIRONMENT", "local")
    monkeypatch.setenv("SESSIONS_PATH", str(tmp_path / "s.json"))
    monkeypatch.setenv("TRAIL_ANCHORS_PATH", str(tmp_path / "a.json"))
    monkeypatch.setattr(session_manager, "SESSIONS_PATH", str(tmp_path / "s.json"))
    scripts = os.path.join(os.path.dirname(__file__), "..", "scripts")
    monkeypatch.syspath_prepend(scripts)
    import demo_walkthrough  # type: ignore[import-not-found]

    demo_walkthrough.run(tmp_path / "out")
    names = sorted(p.name for p in (tmp_path / "out").iterdir())
    assert names == ["oscal.json", "racm.xlsx", "report.md", "working-papers.xlsx"]
    report = (tmp_path / "out" / "report.md").read_text()
    assert "Return for rework (Planning)" in report
    assert "## Engagement Conclusion" in report
