"""
Download formats for audit artifacts: RACM and working papers as .xlsx, the
final report as Markdown, and the report's OSCAL-shaped part as JSON.

Artifact text is model output (and, for the scope, user input), so:

* every spreadsheet cell goes through :func:`sanitize_cell`, which defuses
  spreadsheet formula injection (CSV/DDE injection) and strips characters
  that are illegal in XLSX XML;
* the Markdown report has image links removed (a rendered image link is a
  blind data-exfiltration / SSRF channel).

Each export also states the artifact's review state (e.g. a draft rejected by
QA) so a downloaded file cannot be mistaken for an approved one.
"""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from datetime import UTC, datetime
from typing import Any, Iterable, Optional

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

from swarm.schema import FinalReportSchema, RiskControlMatrixSchema, WorkingPaperSchema

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# A cell starting with one of these is interpreted as a formula by Excel /
# LibreOffice / Sheets (tab and CR as well, per OWASP CSV-injection guidance).
_FORMULA_PREFIXES = ("=", "+", "-", "@", "\t", "\r")

_PHASE_LABELS = {1: "Planning", 2: "Fieldwork", 3: "Reporting"}


def sanitize_cell(value: Any) -> str:
    """Render ``value`` as inert spreadsheet text."""
    text = "" if value is None else str(value)
    text = ILLEGAL_CHARACTERS_RE.sub("", text)
    if text.startswith(_FORMULA_PREFIXES):
        text = "'" + text
    return text


_MD_IMAGE = re.compile(r"!\[[^\]]*\]\([^)]*\)|!\[[^\]]*\]\[[^\]]*\]")
_HTML_IMG = re.compile(r"<img\b[^>]*>", re.IGNORECASE)


def sanitize_report(text: str) -> str:
    """Strip Markdown / HTML images from model-written report text."""
    text = _MD_IMAGE.sub("[Image removed for security]", text)
    return _HTML_IMG.sub("[Image removed for security]", text)


@dataclass(frozen=True)
class ExportContext:
    session_id: str
    session_name: str
    status: str

    def artifact_state(self, phase: int) -> str:
        """Human-readable review state of the phase-``phase`` artifact."""
        status = self.status
        if status == f"QA_REJECTED_PHASE_{phase}":
            return "DRAFT REJECTED BY QA — not approved"
        if status in (f"RUNNING_PHASE_{phase}", f"ERROR_PHASE_{phase}"):
            return "Previous draft — this phase is being re-run or failed"
        if status == f"WAITING_HUMAN_GATE_{phase}":
            return f"Awaiting human approval at Gate {phase}"
        return f"Approved at Gate {phase}"


def _slug(session_id: str) -> str:
    return re.sub(r"[^A-Za-z0-9-]", "", session_id)[:36] or "session"


def export_filename(ctx: ExportContext, stem: str, ext: str) -> str:
    return f"{stem}-{_slug(ctx.session_id)[:8]}.{ext}"


def _write_sheet(wb: Workbook, title: str, headers: list[str], rows: Iterable[list]):
    ws = wb.create_sheet(title)
    ws.append([sanitize_cell(h) for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append([sanitize_cell(v) for v in row])
    ws.freeze_panes = "A2"
    return ws


def _info_rows(ctx: ExportContext, phase: int, notes: list[str]) -> list[list[str]]:
    rows = [
        ["Session", ctx.session_name],
        ["Session ID", ctx.session_id],
        ["Session status", ctx.status],
        [f"{_PHASE_LABELS[phase]} artifact", ctx.artifact_state(phase)],
        ["Exported at (UTC)", datetime.now(UTC).isoformat(timespec="seconds")],
    ]
    rows.extend(["Note", n] for n in notes)
    return rows


def _to_bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _steps(steps: Optional[list[Any]]) -> str:
    return "; ".join(
        f"{s.step_description} (Expect: {s.expected_result})" for s in steps or []
    )


def _yes_no(value: Optional[bool]) -> str:
    return "" if value is None else ("Yes" if value else "No")


def _text(value: Any) -> str:
    return "" if value is None else str(value)


RACM_HEADERS = [
    "Risk ID",
    "Risk Description",
    "Likelihood",
    "Impact",
    "Rating Rationale",
    "Regulatory Mapping",
    "Control ID",
    "Control Description",
    "Control Owner",
    "Frequency",
    "Nature",
    "Type",
    "Key Control",
    "Assertions / Objectives",
    "IPE",
    "ToD Steps",
    "ToE Steps",
    "Substantive Steps",
    "Population Source",
    "Population Completeness",
    "Sample Size",
    "Sampling Method",
    "Period of Reliance",
]


def racm_xlsx(racm: RiskControlMatrixSchema, ctx: ExportContext) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    rows = []
    for risk in racm.risks:
        mapping = ", ".join(risk.regulatory_mapping)
        for c in risk.controls:
            tp = c.testing_procedures
            pop = tp.population
            rows.append(
                [
                    risk.risk_id,
                    risk.description,
                    _text(risk.likelihood),
                    _text(risk.impact),
                    _text(risk.rating_rationale),
                    mapping,
                    c.control_id,
                    c.description,
                    _text(c.control_owner),
                    _text(c.frequency),
                    _text(c.nature),
                    _text(c.control_type),
                    _yes_no(c.key_control),
                    "; ".join(c.assertions),
                    "; ".join(c.ipe),
                    _steps(tp.test_of_design),
                    _steps(tp.test_of_effectiveness),
                    _steps(tp.substantive_testing),
                    pop.source if pop else "",
                    pop.completeness_procedure if pop else "",
                    _text(tp.sample_size),
                    _text(tp.sampling_method),
                    _text(tp.period_of_reliance),
                ]
            )
    _write_sheet(wb, "Controls", RACM_HEADERS, rows)
    _write_sheet(
        wb,
        "Export Info",
        ["Field", "Value"],
        _info_rows(ctx, 1, [f"Theme: {racm.theme}"]),
    )
    return _to_bytes(wb)


WORKING_PAPER_HEADERS = [
    "Control ID",
    "ToD Conclusion",
    "ToE Conclusion",
    "ToE Basis",
    "Items Tested",
    "Exceptions Noted",
    "Result",
    "Preliminary Deficiency",
    "Conclusion",
    "Evidence Quote",
    "Vault ID",
    "Quote Verified in Vault",
    "Legacy Severity",
]


def working_papers_xlsx(papers: WorkingPaperSchema, ctx: ExportContext) -> bytes:
    from swarm.evidence import EvidenceAssuranceProtocol

    wb = Workbook()
    wb.remove(wb.active)  # type: ignore[arg-type]
    rows = []
    for f in papers.findings:
        if f.vault_id_reference and f.exact_quote_from_evidence:
            verified = EvidenceAssuranceProtocol.verify_exact_quote(
                f.vault_id_reference, f.exact_quote_from_evidence
            )
            verified_text = "Yes" if verified else "No"
        else:
            verified_text = "n/a (no evidence)"
        rows.append(
            [
                f.control_id,
                _text(f.tod_conclusion),
                _text(f.toe_conclusion),
                _text(f.toe_basis),
                _text(f.items_tested),
                _text(f.exceptions_noted),
                _text(f.result),
                _yes_no(f.preliminary_deficiency),
                f.test_conclusion,
                f.exact_quote_from_evidence,
                f.vault_id_reference,
                verified_text,
                _text(f.legacy_severity),
            ]
        )
    notes = [
        f"Theme: {papers.theme}",
        "Quote Verified in Vault is re-checked at export time: the quote "
        "must appear verbatim in the stored evidence and the record's "
        "digest must match.",
        "Preliminary Deficiency is a fieldwork flag only. Deficiencies are "
        "classified at engagement level in the report's deficiency evaluation.",
    ]
    if any(f.legacy_severity for f in papers.findings):
        notes.append(
            "Legacy Severity: these findings were recorded before ToD/ToE "
            "conclusions existed; their conclusions were derived from the old "
            "severity label on load."
        )
    _write_sheet(wb, "Findings", WORKING_PAPER_HEADERS, rows)
    _write_sheet(
        wb,
        "Export Info",
        ["Field", "Value"],
        _info_rows(ctx, 2, notes),
    )
    return _to_bytes(wb)


def _one_line(value: Any) -> str:
    return " ".join(str(value or "").split())


def _md_cell(value: Any) -> str:
    return sanitize_report(_one_line(value).replace("|", "\\|"))


def _deficiency_section(report: FinalReportSchema, ctx: ExportContext) -> list[str]:
    scale = _one_line(report.deficiency_scale) or "not stated"
    if ctx.status == "COMPLETED":
        note = (
            "Engagement-level evaluation proposed by the reporting crew; the "
            "report containing it was approved at Gate 3 (see Approval Trail)."
        )
    else:
        note = (
            "Draft engagement-level evaluation proposed by the reporting crew, "
            "for the auditor's judgement at Gate 3."
        )
    lines = [
        "## Deficiency Evaluation (proposed)",
        "",
        f"> {note} Scale: {scale}.",
        "",
    ]
    if not report.deficiency_evaluations:
        lines += ["No deficiencies were proposed.", ""]
        return lines
    lines += [
        "| ID | Title | Findings | Risks | Likelihood | Magnitude "
        "| Classification | Compensating controls | Rationale |",
        "|---|---|---|---|---|---|---|---|---|",
    ]
    for e in report.deficiency_evaluations:
        cells = [
            e.deficiency_id,
            e.title,
            ", ".join(e.related_findings),
            ", ".join(e.related_risks),
            e.likelihood,
            e.magnitude,
            e.classification,
            e.compensating_controls,
            e.rationale,
        ]
        lines.append("| " + " | ".join(_md_cell(c) for c in cells) + " |")
    lines.append("")
    return lines


def report_markdown(
    report: FinalReportSchema, trail: list[dict[str, str]], ctx: ExportContext
) -> str:
    lines = [
        sanitize_report(f"# GRC Audit Report — {_one_line(ctx.session_name)}"),
        "",
        f"> Report status: {ctx.artifact_state(3)} (session status {ctx.status}).",
        "",
        "## Executive Summary",
        "",
        sanitize_report(report.executive_summary),
        "",
        "## Detailed Report",
        "",
        sanitize_report(report.detailed_report),
        "",
        *_deficiency_section(report, ctx),
        "## Approval Trail",
        "",
    ]
    if not trail:
        lines.append("No approvals recorded.")
    for e in trail:
        entry = (
            f"- **{_one_line(e.get('gate'))}** — {_one_line(e.get('human'))} "
            f"at {_one_line(e.get('timestamp'))}"
        )
        if e.get("action"):
            entry += f" ({_one_line(e['action'])})"
        if e.get("reason"):
            entry += f" — reason: {_one_line(e['reason'])}"
        if e.get("qa_rejection_reason"):
            entry += f" — overrode QA: {_one_line(e['qa_rejection_reason'])}"
        lines.append(sanitize_report(entry))
    return "\n".join(lines) + "\n"


def oscal_json(report: FinalReportSchema) -> Optional[bytes]:
    """The report's OSCAL-shaped part, or None when the report has none."""
    if report.oscal_sar is None:
        return None
    return json.dumps(report.oscal_sar.model_dump(mode="json"), indent=2).encode()
